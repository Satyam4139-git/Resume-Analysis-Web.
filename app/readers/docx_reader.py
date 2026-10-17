import logging

from docx import Document


logger = logging.getLogger("ResumeAnalyzer")


class DOCXReader:
    """
    Reads text from DOCX files.

    Reads both normal paragraphs and tables because
    resumes frequently store information inside tables.
    """

    def read(self, file_path):

        document = Document(
            file_path
        )

        text_parts = []

        # ======================================================
        # PARAGRAPHS
        # ======================================================

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                text_parts.append(
                    text
                )

        # ======================================================
        # TABLES
        # ======================================================

        for table in document.tables:

            for row in table.rows:

                row_values = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:

                        row_values.append(
                            cell_text
                        )

                if row_values:

                    text_parts.append(
                        " | ".join(row_values)
                    )

        text = "\n".join(
            text_parts
        )

        logger.info(
            "DOCX text extraction completed. "
            "Characters extracted: %s",
            len(text)
        )

        return text